"""Documentos de salida en Word (y PDF si hay Microsoft Word): propuesta económica y brief para arquitecto.

Todo número sale de la cotización del motor; aquí solo se da formato.
"""
import re
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from app import config

ROJO = RGBColor(0xC6, 0x28, 0x28)
GRIS = RGBColor(0x5F, 0x6B, 0x76)
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]


def _dinero(n):
    return "—" if n is None else f"${n:,.2f}"


def _fecha(iso: str) -> str:
    d = date.fromisoformat(iso)
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def _supuestos() -> dict:
    """Lee los textos de SUPUESTOS.md (única fuente) para citarlos en la propuesta."""
    textos = {}
    ruta = config.RAIZ / "SUPUESTOS.md"
    for linea in ruta.read_text(encoding="utf-8").splitlines():
        m = re.match(r"\|\s*(SUP-\d+)\s*\|\s*(.+?)\s*\|", linea)
        if m:
            textos[m.group(1)] = re.sub(r"\*\*|`", "", m.group(2))
    return textos


def _documento(titulo_marca: str) -> Document:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
        p = s.header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{titulo_marca} · BORRADOR · DATOS SIMULADOS — no es una oferta oficial")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), ROJO
        f = s.footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rf = f.add_run("Generado por el Copiloto Comercial (Hackathon Claude). Precios del catálogo simulado; "
                       "ningún monto fue calculado por la IA.")
        rf.font.size, rf.font.color.rgb = Pt(8), GRIS
    return doc


def _marca_borrador(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BORRADOR · DATOS SIMULADOS")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(20), ROJO


def _tabla(doc, encabezados, filas, anchos=None, tam=9):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(encabezados):
        t.rows[0].cells[i].text = h
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = "" if v is None else str(v)
    for fila in t.rows:
        for i, c in enumerate(fila.cells):
            if anchos:
                c.width = Cm(anchos[i])
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(tam)
    return t


def _pequeno(doc, texto, color=GRIS):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size, r.font.color.rgb = Pt(8.5), color
    return p


# ------------------------------------------------------------------------------- propuesta
def propuesta(sesion, r: dict, impacto: dict) -> str:
    c, t, cli = r["cotizacion"], r["triaje"], sesion.cliente
    doc = _documento("Propuesta económica")
    _marca_borrador(doc)
    doc.add_heading(f"Propuesta económica — {cli['empresa']}", level=1)
    doc.add_paragraph(f"Folio: S-{sesion.id:04d} · Fecha: {_fecha(c['fecha'])} · "
                      f"Vigencia: 30 días naturales (hasta el {_fecha(c['vigencia_hasta'])}) · "
                      f"Plazo del contrato: {c['plazo_meses']} meses")
    doc.add_paragraph(f"Cliente: {cli['empresa']} · {cli['industria']} · {cli['ciudad']} · {cli['empleados']} empleados")
    if getattr(sesion, "vendedor", ""):
        doc.add_paragraph(f"Elaboró: {sesion.vendedor} · Ejecutivo comercial")
    necesidad = next((n for n in r["notas"] if n["tema"] == "Necesidad principal" and n["estado"] != "faltante"), None)
    if necesidad:
        doc.add_paragraph(f"Necesidad: {necesidad['resumen']}.")

    doc.add_heading("Servicios propuestos", level=2)
    filas = []
    for x in c["renglones"]:
        desc = ", ".join(f"{d['concepto']} −{d['pct']}%" for d in x["descuentos"]) or "—"
        filas.append([x["sku"], x["nombre"], "" if x["cantidad"] is None else f"{x['cantidad']:g}", x["unidad"],
                      _dinero(x["mensual_lista"]), desc, _dinero(x["mensual_neto"]), _dinero(x["cargo_unico"]),
                      x["estado"]])
    _tabla(doc, ["SKU", "Servicio", "Cant.", "Unidad", "Mensual lista", "Descuentos", "Mensual neto",
                 "Cargo único", "Estado"], filas, [2.0, 3.2, 1.0, 2.0, 1.8, 2.4, 1.8, 1.8, 1.8], tam=8)

    doc.add_heading("Totales (MXN)", level=2)
    m, u = c["mensual"], c["cargos_unicos"]
    _tabla(doc, ["Concepto", "Subtotal", f"IVA {c['iva_pct']}%", "Total"], [
        ["Mensualidad", _dinero(m["subtotal"]), _dinero(m["iva"]), _dinero(m["total"])],
        ["Cargos únicos", _dinero(u["subtotal"]), _dinero(u["iva"]), _dinero(u["total"])],
        [f"Total del contrato ({c['plazo_meses']} meses)", _dinero(c["total_contrato"]["subtotal"]),
         _dinero(c["total_contrato"]["total"] - c["total_contrato"]["subtotal"]), _dinero(c["total_contrato"]["total"])],
    ], [6, 3.5, 3.5, 3.5], tam=9.5)
    doc.add_paragraph(f"Los cargos únicos se pueden pagar en 3 mensualidades de {_dinero(u['opcion_3_mensualidades'])} "
                      "(IVA incluido).")
    if c["requiere_autorizacion_gerente"]:
        p = doc.add_paragraph()
        rr = p.add_run(f"Requiere autorización del gerente: el descuento total llega a {c['descuento_max_pct']}% "
                       "(el vendedor puede dar hasta 10% sin autorización).")
        rr.bold, rr.font.color.rgb = True, ROJO
    if c["renglones_pendientes_arquitecto"]:
        doc.add_paragraph("Pendiente de arquitecto de soluciones (sin precio en esta propuesta): "
                          + ", ".join(c["renglones_pendientes_arquitecto"]) + ".")
    if c["renglones_por_definir"]:
        doc.add_paragraph("Por definir con datos del cliente (no incluidos en los totales): "
                          + ", ".join(c["renglones_por_definir"]) + ".")

    doc.add_heading("Condiciones", level=2)
    for texto in [
        "Precios en pesos mexicanos (MXN); no incluyen IVA. El IVA (16%) se desglosa arriba.",
        "Plazo de 12 meses: precio de lista; 24 meses: 5% en mensualidades; 36 meses: 10% en mensualidades.",
        "Los cargos únicos no tienen descuento por plazo.",
        "Los servicios según consumo se cobran por lo que realmente se use.",
        "Vigencia de la propuesta: 30 días naturales.",
    ]:
        doc.add_paragraph(texto, style="List Bullet")

    doc.add_heading("Fuente de cada renglón", level=2)
    sup = _supuestos()
    usados = set()
    for x in c["renglones"]:
        notas = f" — {'; '.join(x['notas'])}" if x["notas"] else ""
        cita = f" — Cliente: «{x['cita']}»" if x.get("cita") else (f" — {x['origen']}" if x.get("origen") != "cliente" else "")
        _pequeno(doc, f"{x['sku']}: {x['fuente']}{notas}{cita}")
        usados |= set(x["supuestos"])
    if usados:
        doc.add_heading("Supuestos aplicados (ajustables)", level=3)
        for s in sorted(usados):
            _pequeno(doc, f"{s}: {sup.get(s, '(ver SUPUESTOS.md)')}")
    _pequeno(doc, c["aviso"], ROJO)

    nombre = f"propuesta_S{sesion.id:04d}_{cli['id']}.docx"
    doc.save(config.SALIDAS_DIR / nombre)
    return nombre


# ------------------------------------------------------------------------------- brief
def brief(sesion, r: dict) -> str:
    c, t, cli = r["cotizacion"], r["triaje"], sesion.cliente
    doc = _documento("Brief para arquitecto")
    doc.add_heading(f"Brief para arquitecto — {cli['empresa']}", level=1)
    _pequeno(doc, "Objetivo: que el cliente no tenga que repetir su explicación. Todo lo de abajo viene de la "
                  "llamada (con cita) o del catálogo.", GRIS)
    doc.add_paragraph(f"Ficha: {cli['empresa']} · {cli['industria']} · {cli['ciudad']} · {cli['empleados']} empleados "
                      f"· Folio S-{sesion.id:04d} · {_fecha(c['fecha'])}")
    doc.add_paragraph(f"Mensaje inicial del cliente: «{cli['mensaje_inicial']}»")

    doc.add_heading("Por qué entra un arquitecto", level=2)
    _tabla(doc, ["Regla", "Qué dice", "Frase del cliente que la activó"],
           [[a["regla"], a["texto"], f"«{a['cita']}»" if a["cita"] else (a.get("detalle") or "")]
            for a in t["reglas_activadas"]], [2.5, 7, 7.5], tam=8.5)

    cotizado = [x for x in c["renglones"] if x["estado"] == "cotizado"]
    doc.add_heading("Ya cotizado como estándar", level=2)
    if cotizado:
        for x in cotizado:
            doc.add_paragraph(f"{x['sku']} {x['nombre']} × {x['cantidad']:g} — {_dinero(x['mensual_neto'])}/mes, "
                              f"{_dinero(x['cargo_unico'])} único", style="List Bullet")
        doc.add_paragraph(f"Subtotal estándar: {_dinero(c['mensual']['subtotal'])}/mes + "
                          f"{_dinero(c['cargos_unicos']['subtotal'])} de cargos únicos (sin IVA).")
    else:
        doc.add_paragraph("Nada todavía.")

    doc.add_heading("Pendiente de diseño", level=2)
    pend = [x for x in c["renglones"] if x["estado"] != "cotizado"]
    for x in pend:
        motivo = "; ".join(x["notas"]) or x["estado"]
        doc.add_paragraph(f"{x['sku']} {x['nombre']} — {motivo}", style="List Bullet")
    if not pend:
        doc.add_paragraph("Nada.")

    doc.add_heading("Fuera de catálogo (no se prometió nada)", level=2)
    for x in t["fuera_de_catalogo"]:
        doc.add_paragraph(f"{x['que_pide']} — «{x['cita']}»", style="List Bullet")
    if not t["fuera_de_catalogo"]:
        doc.add_paragraph("Nada.")

    doc.add_heading("Lo que ya dijo el cliente", level=2)
    for n in r["notas"]:
        if n["estado"] != "faltante":
            doc.add_paragraph(f"{n['tema']}: {n['resumen']}" + (f" — «{n['cita']}»" if n.get("cita") else ""),
                              style="List Bullet")

    doc.add_heading("Preguntas abiertas", level=2)
    preguntas = [f["pregunta"] for f in t["faltantes"]] + [f"{n['tema']}: {n['resumen']}" for n in r["notas"]
                                                           if n["estado"] == "faltante"]
    for q in dict.fromkeys(preguntas):
        doc.add_paragraph(q, style="List Bullet")
    if not preguntas:
        doc.add_paragraph("Ninguna.")
    _pequeno(doc, c["aviso"], ROJO)

    nombre = f"brief_arquitecto_S{sesion.id:04d}_{cli['id']}.docx"
    doc.save(config.SALIDAS_DIR / nombre)
    return nombre


def a_pdf(nombre_docx: str) -> str | None:
    """Convierte con Microsoft Word si está instalado (opcional). Si falla, no pasa nada."""
    try:
        import pythoncom
        from docx2pdf import convert
    except ImportError:
        return None
    pdf = nombre_docx.replace(".docx", ".pdf")
    try:
        pythoncom.CoInitialize()
        convert(str(config.SALIDAS_DIR / nombre_docx), str(config.SALIDAS_DIR / pdf))
        return pdf if (config.SALIDAS_DIR / pdf).exists() else None
    except Exception:
        return None
    finally:
        pythoncom.CoUninitialize()


def generar(sesion, r: dict, impacto: dict) -> dict:
    config.SALIDAS_DIR.mkdir(parents=True, exist_ok=True)
    archivos = {"propuesta_docx": propuesta(sesion, r, impacto)}
    if r["triaje"]["estado"] == "rojo":
        archivos["brief_docx"] = brief(sesion, r)
    if config.GENERAR_PDF:
        archivos["propuesta_pdf"] = a_pdf(archivos["propuesta_docx"])
    return {k: v for k, v in archivos.items() if v}
